"""
生成硕士论文中期报告 DOCX。

设计预设：standard_business_brief
命名覆盖：中文字体使用 Microsoft YaHei，避免中文渲染缺字。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MARKDOWN_PATH = ROOT / "docs" / "midterm_report_draft.md"
OUTPUT_PATH = ROOT / "docs" / "midterm_report.docx"

FIGURES = [
    ("图 1 训练损失曲线", ROOT / "outputs" / "M1-baseline" / "viz" / "fig_loss_curves.png"),
    ("图 2 节点类型混淆矩阵", ROOT / "outputs" / "M1-baseline" / "viz" / "fig_cm.png"),
    ("图 3 跨模态对齐热图", ROOT / "outputs" / "M1-baseline" / "viz" / "fig_attention.png"),
    ("图 4 预测树与真实树结构对比", ROOT / "outputs" / "M1-baseline" / "viz" / "fig_tree_compare.png"),
    ("图 5 树深度分布对比", ROOT / "outputs" / "M1-baseline" / "viz" / "fig_depth_hist.png"),
]


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(18, 18, 18)
MUTED = RGBColor(89, 89, 89)
LIGHT_FILL = "F2F4F7"
BORDER = "B7C2D0"


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 italic: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_text(cell, text: str, bold: bool = False, align=None) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=bold, color=INK)


def set_table_borders(table, color: str = BORDER) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = table_widths(cols)
    set_table_widths(table, widths)
    set_cell_margins(table)
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_FILL)
            align = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 16 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, text, bold=(r_idx == 0), align=align)
    doc.add_paragraph()


def table_widths(cols: int) -> list[float]:
    if cols == 2:
        return [2.0, 4.5]
    if cols == 3:
        return [1.65, 2.35, 2.5]
    if cols == 7:
        return [1.0, 0.78, 0.78, 0.75, 0.9, 1.05, 1.24]
    return [6.5 / cols] * cols


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("硕士论文中期报告")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    add_field(footer, "PAGE")
    run = footer.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("硕士论文中期报告")
    set_run_font(r, size=23, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("基于网页图像与源码的通用结构化设计语言生成方法研究")
    set_run_font(r, size=14, color=MUTED)

    rows = [
        ("作者", "赵志鹏"),
        ("专业", "软件工程"),
        ("研究周期", "2025 年 12 月至 2026 年 12 月"),
        ("当前阶段", "论文中期检查，2026 年 6 月"),
        ("定位说明", "Figma JSON 作为实验验证载体，研究目标抽象为通用 Design IR"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [1.35, 5.15])
    set_cell_margins(table)
    set_table_borders(table)
    for i, (label, value) in enumerate(rows):
        set_cell_shading(table.cell(i, 0), LIGHT_FILL)
        set_cell_text(table.cell(i, 0), label, bold=True)
        set_cell_text(table.cell(i, 1), value)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "本报告基于开题方案、当前代码实现和 M1 阶段实验结果整理，"
        "重点说明课题从单一工具格式生成升级为通用结构化设计语言生成后的研究进展。"
    )
    set_run_font(r, size=11, color=INK)


def add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        clean = part.strip("`") if part.startswith("`") and part.endswith("`") else part
        run = p.add_run(clean)
        set_run_font(run, size=10.5 if part.startswith("`") else 11, color=INK)
        if part.startswith("`"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build_doc() -> None:
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    add_title_page(doc)

    lines = MARKDOWN_PATH.read_text(encoding="utf-8").splitlines()
    idx = 0
    in_title_block = True
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line:
            idx += 1
            continue

        if in_title_block:
            if line.startswith("## "):
                in_title_block = False
            else:
                idx += 1
                continue

        if line.startswith("|"):
            rows, idx = parse_markdown_table(lines, idx)
            add_table(doc, rows)
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\d+\. ", line):
            add_rich_paragraph(doc, re.sub(r"^\d+\. ", "", line), style="List Number")
        elif line.startswith("- "):
            add_rich_paragraph(doc, line[2:], style="List Bullet")
        elif line.startswith("# "):
            pass
        else:
            add_rich_paragraph(doc, line)
        idx += 1

    doc.add_page_break()
    doc.add_heading("附录：阶段性实验图表", level=1)
    for caption, path in FIGURES:
        if not path.exists():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
        doc.add_picture(str(path), width=Inches(6.1))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_doc()
